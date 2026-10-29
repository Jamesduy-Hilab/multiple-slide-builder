import sys
import argparse
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def hex_to_rgb(hex_str):
    hex_str = hex_str.lstrip('#')
    return RGBColor(*(int(hex_str[i:i+2], 16) for i in (0, 2, 4)))

def recolor_deck(input_path, output_path=None, bg_hex="FFFFFF", accent_hex="CE2029", text_hex="1E1D1B"):
    if not output_path:
        output_path = input_path
        
    prs = Presentation(input_path)
    
    HILAB_RED = hex_to_rgb(accent_hex)
    CARBON_BLACK = hex_to_rgb(text_hex)
    BG_COLOR = hex_to_rgb(bg_hex)
    LIGHT_GRAY = RGBColor(245, 245, 245)
    
    slide_width = prs.slide_width
    slide_height = prs.slide_height
    
    print(f"Slide Dimensions: {slide_width} x {slide_height}")
    
    for slide_idx, slide in enumerate(prs.slides):
        print(f"Recoloring Slide {slide_idx + 1}/{len(prs.slides)}...")
        
        # 1. Recolor Slide Background Property
        try:
            background = slide.background
            fill = background.fill
            fill.solid()
            fill.fore_color.rgb = BG_COLOR
        except Exception as e:
            pass
            
        # 2. Recolor all shapes on the slide
        for shape in slide.shapes:
            is_bg_shape = False
            if shape.width >= slide_width - Inches(0.1) and shape.height >= slide_height - Inches(0.1):
                is_bg_shape = True
                
            # Recolor shape fill
            if hasattr(shape, "fill") and shape.fill:
                if shape.fill.type == 1:  # Solid fill
                    color = shape.fill.fore_color.rgb if shape.fill.fore_color.type == 1 else None
                    if color:
                        r, g, b = color
                        if is_bg_shape or (r < 50 and g < 50 and b < 50):
                            shape.fill.solid()
                            shape.fill.fore_color.rgb = BG_COLOR
                        elif (r < 150 and g > 150 and b < 100):  # Lime green/greenish
                            shape.fill.solid()
                            shape.fill.fore_color.rgb = HILAB_RED
                    else:
                        if is_bg_shape:
                            shape.fill.solid()
                            shape.fill.fore_color.rgb = BG_COLOR
                elif shape.fill.type is None:
                    if is_bg_shape:
                        shape.fill.solid()
                        shape.fill.fore_color.rgb = BG_COLOR
                    
            # Recolor shape outline/line
            if hasattr(shape, "line") and shape.line:
                try:
                    if shape.line.fill.type == 1:
                        color = shape.line.fill.fore_color.rgb if shape.line.fill.fore_color.type == 1 else None
                        if color:
                            r, g, b = color
                            if r < 150 and g > 150 and b < 100:
                                shape.line.color.rgb = HILAB_RED
                            elif r > 200 and g > 200 and b > 200:
                                shape.line.color.rgb = LIGHT_GRAY
                except Exception:
                    pass
                            
            # Recolor text runs inside the shape
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.font.color and run.font.color.type == 1:
                            color = run.font.color.rgb
                            r, g, b = color
                            # If text is white or light gray, change to dark text
                            if r > 200 and g > 200 and b > 200:
                                run.font.color.rgb = CARBON_BLACK
                            # If text is green/accent, change to Red
                            elif r < 150 and g > 150 and b < 100:
                                run.font.color.rgb = HILAB_RED
                        else:
                            run.font.color.rgb = CARBON_BLACK
                            
    prs.save(output_path)
    print(f"Successfully recolored slides saved to: {output_path}")

def main():
    parser = argparse.ArgumentParser(description="Recolor PPTX Presentation to Light Background & Brand Accent Theme")
    parser.add_argument("--input", "-i", required=True, help="Input PPTX file path")
    parser.add_argument("--output", "-o", help="Output PPTX file path (defaults to overwriting input)")
    parser.add_argument("--bg", default="FFFFFF", help="Background hex color code (default: FFFFFF)")
    parser.add_argument("--accent", default="CE2029", help="Accent hex color code (default: CE2029)")
    parser.add_argument("--text", default="1E1D1B", help="Text hex color code (default: 1E1D1B)")
    
    args = parser.parse_args()
    recolor_deck(args.input, args.output, args.bg, args.accent, args.text)

if __name__ == "__main__":
    main()
